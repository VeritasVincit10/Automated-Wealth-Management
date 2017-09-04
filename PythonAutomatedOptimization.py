import glob
import os
import openpyxl

recdir = r"C:\Users\Mamadou\Documents\FWMshared\Optimization Outputs" #This is the recommendation directory, Change as needed
posdir = r"C:\Users\Mamadou\Downloads"
os.chdir(recdir)
newestrec = max(glob.iglob('*.xlsx'), key=os.path.getmtime) #Make Sure to save most recent recommendation
rec = openpyxl.load_workbook(newestrec)
os.chdir(posdir)
newestpos = max(glob.iglob('*.xlsx'), key=os.path.getmtime)
pos = openpyxl.load_workbook(newestpos)

clws = pos.active
rcws = rec.get_sheet_by_name("Wertz 5.31") #match the client tab name (please enter exactly)
ticsht = rec.get_sheet_by_name("State0new")


posnames = list()
accnames = list()
tickpos = list()
desc = list()
for x in clws.columns[3]:
    posnames.append(x.value)
for x in clws.columns[1]:
    accnames.append(x.value)
for x in clws.columns[6]:
    x = x.value[0:x.value.find('_')]
    desc.append(x)
for x in clws.columns[12]:
    tickpos.append(x.value)

accnames = [x[5:]for x in accnames]
accnames = [int(x) for x in accnames[1:]]
posnames.pop(0)
tickpos.pop(0)
desc.pop(0)
tickpos = [round(100*i,3) for i in tickpos]
posacc = [posnames,accnames,tickpos, desc]
posacc = list(map(list, zip(*posacc)))


nticks = 41 #Enter row number of last ticker here. Can be bigger, but makes code less efficient.
nticks = nticks + 1
acccells = ['G2', 'H2', 'I2', 'J2', 'K2']
accdict = {'G':19, 'H':20, 'I':21, 'J':22, 'K':23}

result = [row[3:41] for row in rcws.columns[6:11]]
for i in result:
    for j in i:
        j.value = None
succ = set()

#Comparison between downloaded files and recommendation files
for x in posacc:
    for y in acccells:
        if x[1] == rcws[y].value:
            for z in range(4, nticks):
                if str(x[0]) in ticsht.cell(row=z, column=1).value:
                    rcws[y[0] + str(z)] = x[2]
                    succ.add(x[0])
                elif rcws.cell(row=z, column=accdict[y[0]]).value is not None and str(x[0]) in rcws.cell(row=z, column=accdict[y[0]]).value:
                    succ.add(x[0])
                    if rcws[y[0] + str(z)].value is None:
                        rcws[y[0] + str(z)] = x[2]
                    else:
                        rcws[y[0] + str(z)] = rcws[y[0] + str(z)].value + x[2]
                elif rcws.cell(row=z, column=accdict[y[0]]).value is not None and len(rcws.cell(row=z, column=accdict[y[0]]).value) > 0:
                    if rcws.cell(row=z, column=accdict[y[0]]).value in str(x[3]):
                        succ.add(x[0])
                        if rcws[y[0] + str(z)].value is None:
                            rcws[y[0] + str(z)] = x[2]
                        else:
                            rcws[y[0] + str(z)] = rcws[y[0] + str(z)].value + x[2]

fails = [x for x in posnames if x not in succ] #Failure (non success) vector

print("The system failed to match the following:",fails)
#This is for my personal testing
colg = [x.value for x in result[0]]
colh = [x.value for x in result[1]]
coli = [x.value for x in result[2]]
#output
os.chdir(recdir) #enter recdir as arg for shared optimization folder
rec.save(filename='test52516.xlsx')#output updated recommendation file. Will save in directory above.

